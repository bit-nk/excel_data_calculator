"""Ledger export tests — CSV content, XLSX readability, status gate."""
from __future__ import annotations

import csv
import io
from datetime import date, datetime, timezone
from decimal import Decimal

import pytest
from openpyxl import load_workbook
from sqlalchemy import select

from app.models.chargeback import ChargebackLineItem, ChargebackReport
from app.models.cloud_accounts import CloudAccount
from app.models.enums import (
    ChargebackStatus,
    CostCategory,
    Platform,
)
from app.services.ledger import CLEARING_ACCOUNT, LEDGER_COLUMNS, LedgerExporter


async def _seed_report(db, business_units, status: ChargebackStatus) -> ChargebackReport:
    db.add(CloudAccount(platform=Platform.AWS.value, account_id="999", account_name="Prod"))

    report = ChargebackReport(
        billing_period="2026-03",
        status=status,
        total_direct_cost=Decimal("100000"),
        total_shared_cost=Decimal("10000"),
        total_management_fee=Decimal("20000"),
        total_cost=Decimal("130000"),
        generated_at=datetime(2026, 4, 1, 12, 0, tzinfo=timezone.utc),
        approved_at=datetime(2026, 4, 2, 12, 0, tzinfo=timezone.utc),
    )
    db.add(report)
    await db.commit()
    await db.refresh(report)

    geq, qst, _ = business_units
    db.add_all([
        ChargebackLineItem(
            report_id=report.id,
            business_unit_id=geq.id,
            platform=Platform.AWS.value,
            category=CostCategory.COMPUTE.value,
            cost_code="GEQ-001",
            direct_cost=Decimal("60000"),
            shared_cost_allocation=Decimal("6000"),
            management_fee=Decimal("12000"),
            total_cost=Decimal("78000"),
        ),
        ChargebackLineItem(
            report_id=report.id,
            business_unit_id=qst.id,
            platform=Platform.MONGODB.value,
            category=CostCategory.DATABASE.value,
            cost_code="QST-002",
            direct_cost=Decimal("40000"),
            shared_cost_allocation=Decimal("4000"),
            management_fee=Decimal("8000"),
            total_cost=Decimal("52000"),
        ),
    ])
    await db.commit()
    return report


@pytest.mark.asyncio
async def test_csv_has_header_and_balanced_clearing(db, business_units):
    report = await _seed_report(db, business_units, ChargebackStatus.APPROVED)
    data = await LedgerExporter(db).to_csv(report.id)
    reader = csv.DictReader(io.StringIO(data.decode("utf-8")))
    rows = list(reader)

    assert reader.fieldnames == LEDGER_COLUMNS
    debit_lines = [r for r in rows if Decimal(r["debit"]) > 0]
    credit_lines = [r for r in rows if Decimal(r["credit"]) > 0]

    assert len(debit_lines) == 2
    assert len(credit_lines) == 1
    assert credit_lines[0]["account_code"] == CLEARING_ACCOUNT

    total_debits = sum(Decimal(r["debit"]) for r in debit_lines)
    total_credits = sum(Decimal(r["credit"]) for r in credit_lines)
    assert total_debits == total_credits == Decimal("130000.00")


@pytest.mark.asyncio
async def test_xlsx_opens_with_expected_columns(db, business_units):
    report = await _seed_report(db, business_units, ChargebackStatus.APPROVED)
    data = await LedgerExporter(db).to_xlsx(report.id)
    wb = load_workbook(io.BytesIO(data))
    ws = wb.active
    header = [c.value for c in ws[1]]
    assert header == LEDGER_COLUMNS


@pytest.mark.asyncio
async def test_cannot_export_unapproved_report(db, business_units):
    report = await _seed_report(db, business_units, ChargebackStatus.CALCULATED)
    with pytest.raises(ValueError, match="APPROVED"):
        await LedgerExporter(db).to_csv(report.id)


@pytest.mark.asyncio
async def test_missing_report_raises(db):
    with pytest.raises(LookupError):
        await LedgerExporter(db).to_csv(99999)
